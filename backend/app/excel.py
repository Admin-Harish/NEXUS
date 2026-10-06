import csv
import io
import json
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

MAX_ROWS = 200

# Recognised header names (lower-cased, punctuation stripped) for each test-case field.
HEADERS = {
    "id": {"id", "test id", "tc id", "case id", "testcase id", "ref", "no", "sno", "s no"},
    "name": {"test case", "testcase", "test", "name", "title", "scenario", "requirement", "description", "test name", "summary"},
    "method": {"method", "http method", "verb"},
    "path": {"path", "endpoint", "url", "route", "api", "uri"},
    "payload": {"body", "payload", "request body", "request", "input", "data"},
    "expected_status": {"expected status", "status", "expected code", "status code", "http status", "expected http status"},
    "expected_contains": {"expected text", "expected contains", "contains", "expected body", "response contains"},
    "expected": {"expected", "expected result", "expected outcome", "result"},
}
METHOD_PATH = re.compile(r"^\s*(GET|POST|PUT|PATCH|DELETE)\s+(/\S*)\s*(.*)$", re.IGNORECASE)


def _norm(header: Any) -> str:
    return re.sub(r"[^a-z0-9 ]", "", str(header or "").lower()).strip()


def read_rows(filename: str, data: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    """Read the first sheet (or a CSV) into (columns, rows) using the first non-empty row as headers."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        table = [list(r) for r in sheet.iter_rows(values_only=True)]
    elif filename.lower().endswith(".csv"):
        table = list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="ignore"))))
    else:
        raise ValueError("Upload an .xlsx or .csv file.")
    table = [r for r in table if any(c not in (None, "") for c in r)]
    if len(table) < 2:
        raise ValueError("The sheet needs a header row and at least one test case row.")
    columns = [str(c).strip() if c not in (None, "") else f"Column {i + 1}" for i, c in enumerate(table[0])]
    rows = []
    for raw in table[1:MAX_ROWS + 1]:
        row = {columns[i]: ("" if v is None else v) for i, v in enumerate(raw[:len(columns)])}
        if any(str(v).strip() for v in row.values()):
            rows.append(row)
    return columns, rows


def _field(row: dict[str, Any], key: str) -> str:
    for column, value in row.items():
        if _norm(column) in HEADERS[key]:
            return str(value).strip()
    return ""


def structured_cases(rows: list[dict[str, Any]]) -> list[dict[str, Any]] | None:
    """Map rows straight to test cases when every row names a method, a path and a status.
    Returns None when the sheet is free-form and needs the LLM to interpret it."""
    cases = []
    for number, row in enumerate(rows, 1):
        method, path = _field(row, "method").upper(), _field(row, "path")
        payload_text = _field(row, "payload")
        combined = METHOD_PATH.match(path) or (METHOD_PATH.match(f"{method} {path}") if method else None)
        if combined:
            method, path = combined.group(1).upper(), combined.group(2)
            payload_text = payload_text or combined.group(3)
        expected = _field(row, "expected")
        status_text = _field(row, "expected_status") or expected
        status = re.search(r"\b([1-5]\d\d)\b", status_text)
        if not (method and path.startswith("/") and status):
            return None
        contains = _field(row, "expected_contains")
        if not contains and expected:
            found = re.search(r"contain(?:s|ing)?\s+[\"“']?([^\"”',]+)", expected, re.IGNORECASE)
            contains = found.group(1).strip() if found else ""
        payload = None
        if payload_text:
            try:
                payload = json.loads(payload_text)
            except json.JSONDecodeError:
                return None
        ref = _field(row, "id") or str(number)
        name = _field(row, "name") or f"{method} {path}"
        cases.append({
            "name": name, "description": f"From spreadsheet row {ref}", "adapter": "api", "priority": "medium",
            "method": method, "path": path, "payload": payload if isinstance(payload, dict) else None,
            "expected_status": int(status.group(1)), "expected_contains": contains or None,
            "steps": [], "evidence": [f"Spreadsheet row {ref}"], "confidence": "high",
            "excel_ref": ref,
        })
    return cases


def export_results(cases: list[dict[str, Any]], results: list[dict[str, Any]], title: str) -> bytes:
    """One row per executed test case, marking which ones NEXUS added to the user's sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Test cases"
    header = ["ID", "Test case", "Source", "Method", "Path", "Body", "Expected status", "Expected text",
              "Result", "Actual status", "Notes"]
    ws.append(header)
    added = 0
    # cases are the executed ones, in run order, so results line up by position
    for index, case in enumerate(cases):
        result = results[index] if index < len(results) else {}
        source = case.get("source", "discovered")
        if source != "excel":
            added += 1
        ws.append([
            case.get("excel_ref") or f"NX-{added}",
            case.get("name"),
            {"excel": "From your sheet", "suggested": "Added by NEXUS", "scenario": "Scenario", "discovered": "Discovered"}.get(source, source),
            case.get("method"), case.get("path"),
            json.dumps(case["payload"]) if case.get("payload") else "",
            case.get("expected_status"), case.get("expected_contains") or "",
            (result.get("status") or "not run").upper(), result.get("actual_status") or "",
            result.get("error") or case.get("description") or "",
        ])
    bold = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor="2457D6")
    fills = {"PASSED": "E7F6EE", "FAILED": "FDECEA", "SKIPPED": "FFF6DD"}
    for row in ws.iter_rows(min_row=2):
        if row[2].value == "Added by NEXUS":
            row[2].font = Font(bold=True, color="2457D6")
        if row[8].value in fills:
            row[8].fill = PatternFill("solid", fgColor=fills[row[8].value])
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, width in zip("ABCDEFGHIJK", (10, 40, 18, 9, 26, 30, 10, 18, 10, 10, 50)):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = "2457D6"
    wb.properties.title = f"NEXUS test cases – {title}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
